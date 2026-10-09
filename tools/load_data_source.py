"""Load capacity_baza_szablon.xlsx plus scenario JSON into capacity.db.

Skips accounts, roles, sessions, admin_settings, and _migrations.
Replaces Excel cells that only contain __FILE__:scenarios/... with the snapshot JSON.
"""

from __future__ import annotations

import json
import re
import sqlite3
import sys
from datetime import datetime, date, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data_source"
EXCEL = SOURCE / "capacity_baza_szablon.xlsx"
DB_PATH = ROOT / "data" / "capacity.db"

AUTH_PROTECTED = {
    "users",
    "user_roles",
    "roles",
    "role_permissions",
    "sessions",
    "password_reset_tokens",
    "password_reset_requests",
}
SKIP = AUTH_PROTECTED | {
    "admin_settings",
    "sqlite_sequence",
    "sqlite_stat1",
    "sqlite_stat4",
    "_migrations",
    "_INSTRUKCJA",
}
IMPORT_ORDER = [
    "working_days",
    "machine_types",
    "process_phases",
    "machines",
    "part_designations",
    "nests",
    "nest_machines",
    "machine_alternatives",
    "projects",
    "project_volumes",
    "project_volumes_contract",
    "project_eop_extensions",
    "project_notes",
    "parts",
    "part_volume_by_year",
    "part_volume_contract_by_year",
    "part_volume_share_by_year",
    "part_volume_contract_share_by_year",
    "operations",
    "operation_volume_by_year",
    "operation_set_members",
    "call_off_comparisons",
    "call_off_volumes",
    "scenarios",
]
DATE_TEXT = {"sop", "eop", "eop_original", "eop_extension", "eop_before", "eop_after"}
IDENT = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
MACHINE_NO = re.compile(r"^\d+(/\d+)?$")
FILE_MARKER = "__FILE__:"


def ident(name: str) -> str:
    if not IDENT.match(name):
        raise SystemExit(f"unsafe name: {name}")
    return name


def pad2(n: int) -> str:
    return f"{n:02d}"


def dotted(d: datetime | date) -> str:
    return f"{pad2(d.day)}.{pad2(d.month)}.{d.year}"


def excel_serial_dotted(value: float) -> str | None:
    whole = int(value)
    if whole < 25000 or whole > 80000:
        return None
    dt = datetime(1970, 1, 1, tzinfo=timezone.utc) + timedelta(days=whole - 25569)
    return dotted(dt)


def normalize_date_text(value: str) -> str:
    text = value.strip()
    iso = re.match(r"^(\d{4})-(\d{2})-(\d{2})", text)
    if iso:
        return f"{iso.group(3)}.{iso.group(2)}.{iso.group(1)}"
    dmy = re.match(r"^(\d{1,2})[./](\d{1,2})[./](\d{4})$", text)
    if dmy:
        return f"{pad2(int(dmy.group(1)))}.{pad2(int(dmy.group(2)))}.{dmy.group(3)}"
    return text


def coerce(column: str, value):
    if column in DATE_TEXT:
        if value is None or value == "":
            return None
        if isinstance(value, datetime):
            return dotted(value)
        if isinstance(value, date):
            return dotted(value)
        if isinstance(value, bool):
            return "1" if value else "0"
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            serial = excel_serial_dotted(float(value))
            return serial if serial else str(value)
        text = normalize_date_text(str(value))
        return text or None

    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return 1 if value else 0
    if isinstance(value, datetime):
        if value.hour == 0 and value.minute == 0 and value.second == 0:
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    text = str(value).strip()
    if text == "":
        return None
    if re.fullmatch(r"-?\d+(\.\d+)?", text):
        return float(text) if "." in text else int(text)
    return text


def scenario_snapshot(value):
    if not isinstance(value, str) or not value.startswith(FILE_MARKER):
        return coerce("snapshot", value)
    rel = value[len(FILE_MARKER):].strip().replace("\\", "/")
    path = SOURCE / rel
    if not path.is_file():
        print(f"missing scenario file: {rel}", file=sys.stderr)
        return value
    parsed = json.loads(path.read_text(encoding="utf-8"))
    snap = parsed.get("snapshot")
    if isinstance(snap, str):
        return snap
    return json.dumps(snap, ensure_ascii=False, separators=(",", ":"))


def table_names(conn: sqlite3.Connection) -> list[str]:
    rows = conn.execute(
        "SELECT name FROM sqlite_master WHERE type = 'table' AND name NOT LIKE 'sqlite_%' ORDER BY name"
    ).fetchall()
    return [row[0] for row in rows]


def import_order(existing: set[str]) -> list[str]:
    ordered = [name for name in IMPORT_ORDER if name in existing and name not in SKIP]
    for name in sorted(existing):
        if name not in ordered and name not in SKIP:
            ordered.append(name)
    return ordered


def columns_of(conn: sqlite3.Connection, table: str) -> list[str]:
    return [row[1] for row in conn.execute(f"PRAGMA table_info({ident(table)})")]


def main() -> None:
    if not EXCEL.is_file():
        raise SystemExit(f"missing workbook: {EXCEL}")
    if not DB_PATH.is_file():
        raise SystemExit(f"missing database: {DB_PATH}")

    print(f"workbook {EXCEL}")
    print(f"database {DB_PATH}")
    wb = load_workbook(EXCEL, read_only=True, data_only=False)
    sheets = {ws.title: ws for ws in wb.worksheets}

    conn = sqlite3.connect(DB_PATH)
    conn.execute("PRAGMA foreign_keys = OFF")
    existing = set(table_names(conn))
    order = import_order(existing)
    print("tables", len(order))

    def sheet_for(table: str):
        if table in sheets:
            return sheets[table]
        return sheets.get(table[:31])

    conn.execute("BEGIN")
    try:
        for table in reversed(order):
            conn.execute(f"DELETE FROM {ident(table)}")

        counts: dict[str, int] = {}
        for table in order:
            ws = sheet_for(table)
            db_cols = columns_of(conn, table)
            if ws is None:
                counts[table] = 0
                print(f"{table}: no sheet")
                continue
            rows_iter = ws.iter_rows(values_only=True)
            try:
                header = next(rows_iter)
            except StopIteration:
                counts[table] = 0
                print(f"{table}: empty")
                continue
            names = [str(cell).strip() if cell is not None else "" for cell in header]
            cols = [name for name in names if name in db_cols]
            if not cols:
                counts[table] = 0
                print(f"{table}: no matching columns")
                continue
            indexes = [names.index(name) for name in cols]
            sql = (
                f"INSERT INTO {ident(table)} ({', '.join(ident(c) for c in cols)}) "
                f"VALUES ({', '.join('?' for _ in cols)})"
            )
            batch = []
            inserted = 0
            seen_numbers: dict[str, int] = {}
            for raw in rows_iter:
                values = []
                for col, idx in zip(cols, indexes):
                    cell = raw[idx] if idx < len(raw) else None
                    if table == "scenarios" and col == "snapshot":
                        values.append(scenario_snapshot(cell))
                    else:
                        values.append(coerce(col, cell))
                if all(value is None for value in values):
                    continue
                if table == "machines" and "internal_number" in cols:
                    number = values[cols.index("internal_number")]
                    key = str(int(number)) if isinstance(number, float) and number.is_integer() else str(number or "").strip()
                    if MACHINE_NO.match(key):
                        seen_numbers[key] = seen_numbers.get(key, 0) + 1
                batch.append(values)
                if len(batch) >= 500:
                    conn.executemany(sql, batch)
                    inserted += len(batch)
                    batch.clear()
            if batch:
                conn.executemany(sql, batch)
                inserted += len(batch)
            if table == "machines":
                dups = sorted(key for key, count in seen_numbers.items() if count > 1)
                if dups:
                    raise RuntimeError("duplicate internal_number: " + ", ".join(dups))
            counts[table] = inserted
            print(f"{table}: {inserted}")

        conn.execute("COMMIT")
    except Exception:
        conn.execute("ROLLBACK")
        raise
    finally:
        conn.execute("PRAGMA foreign_keys = ON")
        wb.close()

    admin = conn.execute("SELECT username FROM users").fetchall()
    migrations = conn.execute("SELECT COUNT(*) FROM _migrations").fetchone()[0]
    print("users", [row[0] for row in admin])
    print("migrations", migrations)
    for table in ("machines", "projects", "parts", "operations", "scenarios"):
        if table in existing:
            count = conn.execute(f"SELECT COUNT(*) FROM {ident(table)}").fetchone()[0]
            print(f"check {table}", count)
    marker = conn.execute(
        "SELECT COUNT(*) FROM scenarios WHERE snapshot LIKE '__FILE__:%'"
    ).fetchone()[0]
    print("scenario markers left", marker)
    conn.close()


if __name__ == "__main__":
    main()
